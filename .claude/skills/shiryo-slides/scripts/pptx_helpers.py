"""資料作成キット：PowerPoint 版の部品（テーマ E の色。他テーマは色トークンを差し替える）

使い方: from pptx_helpers import *  →  s = new_slide(note) で 1 枚追加（note = 発表者メモ。必須）
座標は 1280x720 の px 指定。最後に check_notes(prs) で全枚にメモがあるか検査してから prs.save()
"""
import math, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# theme E tokens
BG = "FAF4EF"; SURF = "FFFFFF"; INK = "2B2320"; SUB = "6E5F57"; DIM = "8A7D74"; LINE = "E5D8CE"
ACC = "EE5D43"; ACCD = "C24631"; TINT = "FBE3DC"; ONDARK = "FAF4EF"; ACC_ON_DARK = "FF8A72"
EA = "游ゴシック"; LAT = "Arial"

def E(px): return Emu(int(round(px * 9525)))
def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = E(1280), E(720)
blank = prs.slide_layouts[6]

PAGE = [0]
def new_slide(note):
    """1 枚追加。note（発表者メモ）は必須。workflows/speaker-notes.md の型で書く"""
    if not note or not note.strip():
        raise ValueError("発表者メモ（note）が空です。workflows/speaker-notes.md の型で書いてください")
    PAGE[0] += 1; i = PAGE[0]
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    s.notes_slide.notes_text_frame.text = note.strip("\n")
    if i != 1:
        text(s, 1100, 668, 100, 24, str(i), 14, DIM, bold=True, align="r", latin=LAT)
    return s

REQUIRED = ("【時間】", "【このページで伝えること】", "【台本】", "【つなぎ】")
def check_notes(p=None):
    """全スライドに発表者メモがあり、必須の見出しがそろっているか検査する。足りなければ例外"""
    p = p or prs
    bad = []
    for k, s in enumerate(p.slides, 1):
        t = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        miss = [h for h in REQUIRED if h not in t]
        if k == len(p.slides): miss = [h for h in miss if h != "【つなぎ】"]
        if miss: bad.append(f"{k}枚目: {'・'.join(miss)} が無い")
    if bad: raise ValueError("発表者メモが不足:\n" + "\n".join(bad))
    return True

def set_font(run, px, color, bold, latin):
    f = run.font
    f.size = Pt(px * 0.75); f.bold = bold; f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", EA), ("a:cs", EA)):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)

def text(s, x, y, w, h, paras, px, color=INK, bold=False, align="l", anchor="t", ls=1.3, latin=LAT, space=0):
    """paras: str (\n = new paragraph) or list of paragraphs; paragraph = str or list of (txt, {opts})."""
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for c in list(bodyPr): bodyPr.remove(c)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    if isinstance(paras, str): paras = paras.split("\n")
    for k, p in enumerate(paras):
        para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        para.line_spacing = ls
        if space and k: para.space_before = Pt(space)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run(); r.text = t
            set_font(r, o.get("px", px), o.get("color", color), o.get("bold", bold), o.get("latin", latin))
    return tb

def shape(s, kind, x, y, w, h, fill=None, line=None, lw=1.5, radius=None):
    sh = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw * 0.75)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius / min(w, h)
    sh.text_frame.text = ""
    return sh

def rect(s, x, y, w, h, fill=None, line=None, lw=1.5, r=4):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, lw, r)

def circle(s, cx, cy, r, fill=None, line=None, lw=1.5):
    return shape(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill, line, lw)

def style_line(ln, color, w, dash=None, head=False, round_=True):
    spPr = ln._element.spPr
    l = spPr.find(qn("a:ln"))
    if l is None: l = etree.SubElement(spPr, qn("a:ln"))
    l.set("w", str(int(w * 0.75 * 12700)))
    if round_: l.set("cap", "rnd")
    for c in list(l): l.remove(c)
    sf = etree.SubElement(l, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
    if dash:
        d = etree.SubElement(l, qn("a:prstDash")); d.set("val", dash)
    etree.SubElement(l, qn("a:round"))
    if head:
        t = etree.SubElement(l, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")

def line(s, x1, y1, x2, y2, color=INK, w=2, dash=None, head=False):
    ln = s.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))
    style_line(ln, color, w, dash, head)
    return ln

def poly(s, pts, color=INK, w=3, dash=None, head=False, fill=None, closed=False):
    fb = s.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=closed)
    sh = fb.convert_to_shape()
    sh.shadow.inherit = False
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if color: style_line(sh, color, w, dash, head)
    else: sh.line.fill.background()
    return sh

def title(s, t, size=40):
    text(s, 80, 52, 1120, 64, t, size, INK, bold=True, anchor="m", ls=1.1)

def wander(x0, y0, x1, y1, amp=80, waves=2.5, n=160):
    pts = []
    for k in range(n + 1):
        t = k / n
        env = math.sin(math.pi * t)
        x = x0 + (x1 - x0) * t + 55 * env * math.sin(2 * math.pi * waves * t * 1.0 + 0.6)
        y = y0 + (y1 - y0) * t - amp * env * math.sin(2 * math.pi * waves * t * 0.8)
        pts.append((x, y))
    return pts

def path_motif(s, x0, y0, x1, y1, amp=80, labels=True):
    line(s, x0, y0, x1, y1, "D9C7BA", 2.5, dash="dash")
    poly(s, wander(x0, y0, x1, y1, amp), ACC, 6)
    circle(s, x0, y0, 11, INK)
    circle(s, x1, y1, 15, SURF, ACCD, 5)
    circle(s, x1, y1, 5, ACCD)
    if labels:
        mx, my = (x0 + x1) / 2, (y0 + y1) / 2
        text(s, x0 - 30, y0 + 20, 120, 24, "いま", 15, SUB, bold=True)
        text(s, x1 - 80, y1 + 24, 160, 24, "目標", 15, SUB, bold=True, align="c")

